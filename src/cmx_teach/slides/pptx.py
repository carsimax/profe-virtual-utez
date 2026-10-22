"""Extrae texto, notas e imágenes de un .pptx a JSON para rehacerlo como deck HTML.
Uso: uv run slides-pptx entrada.pptx [salida_dir]"""
import json, sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def main():
    src, out = Path(sys.argv[1]), Path(sys.argv[2] if len(sys.argv) > 2 else "pptx-extract")
    (out / "assets").mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(src))
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        item = {"n": n, "title": title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "",
                "blocks": [], "tables": [], "images": [], "notes": ""}
        for sh in walk(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.has_text_frame:
                paras = [{"text": p.text.strip(), "level": p.level} for p in sh.text_frame.paragraphs if p.text.strip()]
                if paras:
                    item["blocks"].append(paras)
            if getattr(sh, "has_table", False) and sh.has_table:
                item["tables"].append([[c.text.strip() for c in r.cells] for r in sh.table.rows])
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = sh.image
                name = f"slide{n:02d}-img{len(item["images"]) + 1}.{img.ext}"
                (out / "assets" / name).write_bytes(img.blob)
                item["images"].append({"file": f"assets/{name}", "w": sh.width, "h": sh.height})
        if slide.has_notes_slide:
            item["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(item)
    (out / "extracted.json").write_text(
        json.dumps({"source": src.name, "count": len(slides), "slides": slides}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(slides)} slides → {out}/extracted.json · imágenes en {out}/assets/")


if __name__ == "__main__":
    main()
