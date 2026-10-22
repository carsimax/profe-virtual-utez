import anydoc
import pytest
from pptx import Presentation
from PIL import Image

from cmx_teach.convert.to_md import main


def _pptx(path, text):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = text
    prs.save(path)


def test_converts_next_to_source(tmp_path, capsys):
    csv = tmp_path / "a.csv"
    csv.write_text("nombre,horas\nPOO,45\n", encoding="utf-8")
    pptx = tmp_path / "b.pptx"
    _pptx(pptx, "Resultado de aprendizaje")
    assert main([str(csv), str(pptx)]) == 0
    assert "POO" in (tmp_path / "a.md").read_text(encoding="utf-8")
    assert "Resultado de aprendizaje" in (tmp_path / "b.md").read_text(encoding="utf-8")
    assert str(tmp_path / "a.md") in capsys.readouterr().out


def test_out_dir_is_created(tmp_path):
    csv = tmp_path / "a.csv"
    csv.write_text("x,y\n1,2\n", encoding="utf-8")
    out = tmp_path / "salida" / "md"
    assert main([str(csv), "--out-dir", str(out)]) == 0
    assert (out / "a.md").exists()
    assert not (tmp_path / "a.md").exists()


def test_failure_does_not_stop_others(tmp_path, capsys):
    good = tmp_path / "a.csv"
    good.write_text("x,y\n1,2\n", encoding="utf-8")
    missing = tmp_path / "nope.docx"
    assert main([str(missing), str(good)]) == 1
    assert (tmp_path / "a.md").exists()
    assert f"docs-to-md: {missing}:" in capsys.readouterr().err


def test_needs_ocr_exits_3(tmp_path, capsys, monkeypatch):
    pdf = tmp_path / "scan.pdf"
    Image.new("RGB", (200, 200), "white").save(pdf, "PDF")
    orig = anydoc.to_markdown

    def fake(path, ocr=None):
        try:
            return orig(path, ocr=ocr)
        except anydoc.NeedsOcrError:
            raise
        except anydoc.ConvertError:
            raise anydoc.NeedsOcrError("scanned")

    monkeypatch.setattr(anydoc, "to_markdown", fake)
    assert main([str(pdf)]) == 3
    assert "--ocr hosted" in capsys.readouterr().err
