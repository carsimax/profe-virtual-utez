import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from cmx_teach.docs.qa import main
from conftest import run_main

FIX = Path(__file__).parent / "fixtures/qa"


def _lo_works():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    return bool(exe) and subprocess.run([exe, "--version"], capture_output=True).returncode == 0


needs_lo = pytest.mark.skipif(not _lo_works(), reason="LibreOffice missing")
OUTLINE = {"cover": "institutional", "hasIndex": True, "headings": [{"key": "Introducción"}, {"key": "Propósito"}], "captions": []}


def test_golden_structure_ok(monkeypatch, capsys):
    assert run_main(main, [FIX / "oficio.docx", "--no-render"], monkeypatch) == 0
    out = capsys.readouterr().out
    assert "títulos" in out and out.rstrip().endswith("OK")


def test_pending_data_listed(tmp_path, monkeypatch, capsys):
    d = Document()
    d.add_heading("Título", 1)
    d.add_paragraph("Folio: [DATO PENDIENTE]")
    f = tmp_path / "p.docx"
    d.save(f)
    run_main(main, [f, "--no-render"], monkeypatch)
    out = capsys.readouterr().out
    assert "DATOS PENDIENTES" in out and "[DATO PENDIENTE]" in out


def test_heading_jump_is_problem(tmp_path, monkeypatch, capsys):
    d = Document()
    d.add_heading("A", 1)
    d.add_heading("B", 3)
    f = tmp_path / "j.docx"
    d.save(f)
    assert run_main(main, [f, "--no-render"], monkeypatch) == 1
    assert "Salto de nivel" in capsys.readouterr().out


def test_malformed_docx_errors(tmp_path, monkeypatch):
    f = tmp_path / "x.docx"
    f.write_bytes(b"not a zip")
    assert run_main(main, [f, "--no-render"], monkeypatch) not in (0, None)


def test_no_args_exit_2(monkeypatch):
    assert run_main(main, [], monkeypatch) == 2


def _check_render(f, capsys, code):
    out = f.parent / "qa" / f.stem
    assert code == 0, capsys.readouterr().out
    assert (out / f"{f.stem}.pdf").exists() and (out / "page-1.png").exists() and (out / "sheet.png").exists()
    pm = json.loads(f.with_suffix(".pages.json").read_text(encoding="utf-8"))
    assert pm["headings"] == {"Introducción": 1, "Propósito": 1} and pm["captions"] == {}
    assert "CAMBIARON" in capsys.readouterr().out


def _prep(tmp_path):
    f = tmp_path / "srs.docx"
    shutil.copy(FIX / "srs.docx", f)
    f.with_suffix(".outline.json").write_text(json.dumps(OUTLINE), encoding="utf-8")
    return f


def test_render_pipeline_with_fake_soffice(tmp_path, monkeypatch, capsys):
    """Fake soffice copies a golden PDF; poppler/Pillow/page-map run for real."""
    bindir = tmp_path / "bin"
    bindir.mkdir()
    fake = bindir / "soffice"
    fake.write_text(f'#!/bin/sh\nfor a; do last="$a"; done\nwhile [ "$1" != "--outdir" ]; do shift; done\ncp "{FIX / "srs.pdf"}" "$2/srs.pdf"\n')
    fake.chmod(0o755)
    monkeypatch.setenv("PATH", f"{bindir}:{os.environ['PATH']}")
    f = _prep(tmp_path)
    _check_render(f, capsys, run_main(main, [f], monkeypatch))
    # second run: same map, indices up to date
    run_main(main, [f], monkeypatch)
    assert "Sin cambios" in capsys.readouterr().out


@needs_lo
def test_render_with_real_libreoffice(tmp_path, monkeypatch, capsys):
    f = _prep(tmp_path)
    code = run_main(main, [f], monkeypatch)
    out = tmp_path / "qa" / "srs"
    assert code == 0, capsys.readouterr().out
    assert (out / "srs.pdf").exists() and (out / "sheet.png").exists()
    assert f.with_suffix(".pages.json").exists()


def test_confirmar_marker_listed(tmp_path, monkeypatch, capsys):
    d = Document()
    d.add_heading("Título", 1)
    d.add_paragraph("Clave de la asignatura: [CONFIRMAR]")
    f = tmp_path / "c.docx"
    d.save(f)
    run_main(main, [f, "--no-render"], monkeypatch)
    assert "[CONFIRMAR]" in capsys.readouterr().out


def test_sidecars_and_qa_dir_live_in_fuentes(tmp_path, monkeypatch, capsys):
    week = tmp_path / "semana-07"
    (week / "fuentes").mkdir(parents=True)
    f = week / "srs.docx"
    shutil.copy(FIX / "srs.docx", f)
    (week / "fuentes" / "srs.outline.json").write_text(json.dumps(OUTLINE), encoding="utf-8")
    run_main(main, [f, "--no-render"], monkeypatch)
    from cmx_teach.docs.paths import qa_dir, sidecar
    assert sidecar(f, ".outline.json") == week / "fuentes" / "srs.outline.json"
    assert qa_dir(f) == week / "fuentes" / "qa" / "srs"
