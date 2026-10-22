"""docs-build: golden specs built through the public CLI and inspected with python-docx/lxml."""

import json
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document

from cmx_teach.docs.build import main
from conftest import run_main

FIX = Path(__file__).parent / "fixtures" / "docs" / "build"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SPECS = ["apa", "examen", "oficio", "planeacion", "srs"]


def build(tmp_path, monkeypatch, name, *extra):
    work = tmp_path / "w"
    if not work.exists():
        shutil.copytree(FIX, work)
    code = run_main(main, [work / f"{name}.json", *extra], monkeypatch)
    return code, work


def xml(path, part="word/document.xml"):
    return zipfile.ZipFile(path).read(part).decode()


@pytest.mark.parametrize("name", SPECS)
def test_builds_and_opens(tmp_path, monkeypatch, name):
    code, work = build(tmp_path, monkeypatch, name)
    assert code == 0
    assert (work / f"{name}.docx").exists()
    assert (work / f"{name}.outline.json").exists()
    assert len(Document(work / f"{name}.docx").paragraphs) > 3


def test_headings_use_styles_with_outline_levels(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    d = Document(work / "srs.docx")
    names = {p.style.name for p in d.paragraphs}
    assert {"Heading 1", "Heading 2"} <= names
    styles = xml(work / "srs.docx", "word/styles.xml")
    for i, lvl in ((1, 0), (2, 1), (3, 2)):
        m = re.search(rf'w:styleId="Heading{i}".*?</w:style>', styles, re.S)
        assert f'<w:outlineLvl w:val="{lvl}"/>' in m.group(0)


def test_numbered_headings_and_lists_use_numpr(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    d = Document(work / "srs.docx")
    h = [p for p in d.paragraphs if p.style.name == "Heading 1"]
    assert all(p._p.pPr.find(f"{W}numPr") is not None for p in h)
    bullets = [p for p in d.paragraphs if p.text in ("Reservas de laboratorio", "Notificaciones")]
    assert bullets and all(p._p.pPr.find(f"{W}numPr") is not None for p in bullets)
    assert not any(p.text.startswith(("•", "- ")) for p in d.paragraphs)


def test_each_list_restarts_with_its_own_num(tmp_path, monkeypatch):
    (tmp_path / "l.json").write_text(json.dumps({"body": ["1. a\n2. b\n  - x", "texto", "1. c\n2. d", {"type": "checklist", "items": ["z"]}]}))
    assert run_main(main, [tmp_path / "l.json"], monkeypatch) == 0
    x = xml(tmp_path / "l.docx")
    used = set(re.findall(r'<w:numId w:val="(\d+)"/>', x))
    assert len(used) >= 3
    numbering = xml(tmp_path / "l.docx", "word/numbering.xml")
    assert all(f'<w:num w:numId="{i}"' in numbering for i in used)
    assert 'w:numFmt w:val="decimal"' in numbering and 'w:numFmt w:val="bullet"' in numbering


@pytest.mark.parametrize("name", ["srs", "planeacion", "apa", "examen"])
def test_tables_have_dxa_grid_summing_to_text_width(tmp_path, monkeypatch, name):
    _, work = build(tmp_path, monkeypatch, name)
    d = Document(work / f"{name}.docx")
    sect = d.sections[-1]
    usable = sect.page_width.twips - sect.left_margin.twips - sect.right_margin.twips
    assert d.tables
    top = [t for t in d.tables if t._tbl.getparent().tag == f"{W}body"]
    checked = 0
    for t in top:
        grid = [int(g.get(f"{W}w")) for g in t._tbl.tblGrid]
        tblw = t._tbl.tblPr.find(f"{W}tblW")
        assert tblw.get(f"{W}type") == "dxa"
        assert sum(grid) == int(tblw.get(f"{W}w")) <= usable
        for c in t._tbl.iter(f"{W}tcW"):
            assert c.get(f"{W}type") == "dxa"
        checked += 1
    assert checked


def test_full_width_tables_equal_usable_width(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    d = Document(work / "srs.docx")
    s = d.sections[-1]
    usable = s.page_width.twips - s.left_margin.twips - s.right_margin.twips
    t = [t for t in d.tables if t.rows[0].cells[0].text == "ID"][0]
    assert sum(int(g.get(f"{W}w")) for g in t._tbl.tblGrid) == usable
    # repeated header row
    assert t._tbl.tr_lst[0].find(f"{W}trPr").find(f"{W}tblHeader") is not None


def test_captions_have_seq_fields(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "planeacion")
    x = xml(work / "planeacion.docx")
    assert re.search(r'w:instr="SEQ (Tabla|Figura) \\\* ARABIC"', x)


def test_toc_cached_entries_and_updatefields(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    x = xml(work / "srs.docx")
    assert 'TOC \\h \\o "1-2"' in x
    assert "1. Introducción" in x and "1.1 Propósito" in x
    assert "<w:updateFields" in xml(work / "srs.docx", "word/settings.xml")
    outline = json.loads((work / "srs.outline.json").read_text())
    assert outline["hasIndex"] and outline["headings"][0]["key"] == "Introducción"


def test_second_pass_fills_page_numbers(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    (work / "srs.pages.json").write_text(
        json.dumps({"headings": {"Introducción": 1, "Requisitos funcionales": 7}, "captions": {}}))
    assert run_main(main, [work / "srs.json"], monkeypatch) == 0
    x = xml(work / "srs.docx")
    assert re.search(r"2\. Requisitos funcionales</w:t><w:tab/><w:t[^>]*>7<", x)
    assert "<w:updateFields" not in xml(work / "srs.docx", "word/settings.xml")


def test_sections_roman_front_and_arabic_body(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    x = xml(work / "srs.docx")
    assert 'w:fmt="lowerRoman"' in x and 'w:fmt="decimal"' in x


def test_footnotes_part_present_for_inline_notes(tmp_path, monkeypatch):
    code, work = build(tmp_path, monkeypatch, "apa")
    assert code == 0
    z = zipfile.ZipFile(work / "apa.docx")
    assert "word/footnotes.xml" in z.namelist()
    assert "<w:footnoteReference" in z.read("word/document.xml").decode()
    assert "<w:footnote " in z.read("word/footnotes.xml").decode()
    assert "footnotes+xml" in z.read("[Content_Types].xml").decode()


def test_inline_syntax(tmp_path, monkeypatch):
    work = tmp_path
    (work / "t.json").write_text(json.dumps({"body": [
        "Hola **negrita** y *cursiva* con `cod` y ==marca== y [enlace](https://x.mx/a)  \nsegunda línea^[nota]"]}))
    assert run_main(main, [work / "t.json"], monkeypatch) == 0
    x = xml(work / "t.docx")
    assert "<w:b/>" in x and "<w:i/>" in x and '<w:highlight w:val="yellow"/>' in x
    assert "<w:hyperlink" in x and "<w:br/>" in x and "<w:footnoteReference" in x
    assert 'w:val="Hyperlink"' in x


def test_apa_preset_cover_header_and_caption(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "apa")
    z = zipfile.ZipFile(work / "apa.docx")
    hdrs = [n for n in z.namelist() if n.startswith("word/header")]
    assert hdrs and any("PAGE" in z.read(n).decode() for n in hdrs)
    assert "Times New Roman" in z.read("word/styles.xml").decode()


def test_figure_png_embedded(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "planeacion")
    z = zipfile.ZipFile(work / "planeacion.docx")
    assert any(n.startswith("word/media/") for n in z.namelist())
    assert "<wp:inline" in z.read("word/document.xml").decode()


def test_exam_questions_and_clave(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "examen")
    plain = xml(work / "examen.docx")
    assert "CLAVE DE RESPUESTAS" not in plain and "← correcta" not in plain and "( X )" not in plain
    assert run_main(main, [work / "examen.json", "--clave"], monkeypatch) == 0
    key = work / "examen_clave.docx"
    assert key.exists()
    x = xml(key)
    assert "CLAVE DE RESPUESTAS" in x and "← correcta" in x and "( X )" in x


def test_default_output_and_explicit_output(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "oficio")
    out = tmp_path / "otro.docx"
    assert run_main(main, [work / "oficio.json", out], monkeypatch) == 0
    assert out.exists() and (tmp_path / "otro.outline.json").exists()


def test_usage_error_exits_2(monkeypatch, capsys):
    assert run_main(main, [], monkeypatch) == 2
    assert "Uso:" in capsys.readouterr().err


def test_unknown_preset_exits_1(tmp_path, monkeypatch, capsys):
    (tmp_path / "x.json").write_text(json.dumps({"preset": "nope", "body": ["hola"]}))
    assert run_main(main, [tmp_path / "x.json"], monkeypatch) == 1
    assert "Preset desconocido" in capsys.readouterr().err


def test_unknown_block_names_the_block(tmp_path, monkeypatch, capsys):
    (tmp_path / "x.json").write_text(json.dumps({"body": [{"type": "zzz"}]}))
    assert run_main(main, [tmp_path / "x.json"], monkeypatch) == 1
    assert "Bloque desconocido" in capsys.readouterr().err


def test_missing_figure_png_hints_docs_visuals(tmp_path, monkeypatch, capsys):
    work = tmp_path / "w"
    shutil.copytree(FIX, work)
    (work / "figs" / "herencia.png").unlink()
    assert run_main(main, [work / "examen.json"], monkeypatch) == 1
    err = capsys.readouterr().err
    assert "Falta herencia.png" in err and "uv run docs-visuals" in err and "--preset" in err


def test_pending_data_passes_through(tmp_path, monkeypatch):
    (tmp_path / "x.json").write_text(json.dumps({"body": ["Fecha: [DATO PENDIENTE]"]}))
    assert run_main(main, [tmp_path / "x.json"], monkeypatch) == 0
    assert "[DATO PENDIENTE]" in xml(tmp_path / "x.docx")


def test_output_is_deterministic(tmp_path, monkeypatch):
    _, work = build(tmp_path, monkeypatch, "srs")
    a = (work / "srs.docx").read_bytes()
    assert run_main(main, [work / "srs.json"], monkeypatch) == 0
    assert (work / "srs.docx").read_bytes() == a


def test_logo_cover_landscape_a4_and_theme(tmp_path, monkeypatch):
    shutil.copy(FIX / "figs" / "avance.png", tmp_path / "logo.png")
    (tmp_path / "x.json").write_text(json.dumps({
        "preset": "corporativo", "page": "a4", "orientation": "landscape", "cover": True, "theme": {"accent": "8B1E3F"},
        "meta": {"title": "T", "institution": "Inst", "logo": "logo.png", "place": "Morelos", "date": "hoy"},
        "front": [{"type": "lot"}, {"type": "lof"}],
        "body": ["# A", {"type": "figure", "src": "logo.png", "caption": "Fig", "width": 0.5},
                 {"type": "table", "headers": ["a", "b"], "rows": [["1", "2"]], "caption": "Tab"},
                 {"type": "signatures", "people": [{"name": "N", "role": "R"}]}]}))
    assert run_main(main, [tmp_path / "x.json"], monkeypatch) == 0
    z = zipfile.ZipFile(tmp_path / "x.docx")
    doc = z.read("word/document.xml").decode()
    assert '<w:pgSz w:w="16838" w:h="11906" w:orient="landscape"/>' in doc
    assert "<w:framePr" in doc and "8B1E3F" in z.read("word/styles.xml").decode()
    assert any("<w:drawing>" in z.read(n).decode() for n in z.namelist() if n.startswith("word/header"))


def test_cover_accent_rule_is_thick_and_styles_are_unique(tmp_path, monkeypatch):
    code, work = build(tmp_path, monkeypatch, "planeacion")
    assert code == 0
    doc = xml(work / "planeacion.docx")
    assert re.search(r'<w:pBdr><w:bottom w:val="single" w:sz="24" w:space="1" w:color="009475"/></w:pBdr>', doc)
    ids = re.findall(r'w:styleId="([^"]+)"', xml(work / "planeacion.docx", "word/styles.xml"))
    assert len(ids) == len(set(ids))


def test_malformed_block_reports_without_traceback(tmp_path, monkeypatch, capsys):
    spec = tmp_path / "bad.json"
    spec.write_text(json.dumps({"preset": "utez", "body": [{"type": "table", "headers": ["a"]}]}), encoding="utf-8")
    assert run_main(main, [spec], monkeypatch) == 1
    err = capsys.readouterr().err
    assert err.startswith("✘") and "Traceback" not in err


def test_spec_in_fuentes_builds_docx_in_parent_and_sidecars_in_fuentes(tmp_path, monkeypatch):
    week = tmp_path / "semana-07"
    src = week / "fuentes"
    shutil.copytree(FIX, src)
    code = run_main(main, [src / "oficio.json"], monkeypatch)
    assert code == 0
    assert (week / "oficio.docx").exists()
    assert not (src / "oficio.docx").exists()
    assert (src / "oficio.outline.json").exists()
    assert not (week / "oficio.outline.json").exists()


def test_spec_in_fuentes_clave_builds_in_parent(tmp_path, monkeypatch):
    week = tmp_path / "semana-07"
    src = week / "fuentes"
    shutil.copytree(FIX, src)
    code = run_main(main, [src / "examen.json", "--clave"], monkeypatch)
    assert code == 0
    assert (week / "examen_clave.docx").exists()
    assert (src / "examen_clave.outline.json").exists()
