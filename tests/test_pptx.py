import json

from PIL import Image
from pptx import Presentation

from cmx_teach.slides.pptx import main
from conftest import run_main


def test_extract(tmp_path, monkeypatch, capsys):
    Image.new("RGB", (4, 4), "red").save(tmp_path / "p.png")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text_frame.text = "Hola"
    s.placeholders[1].text_frame.text = "punto uno"
    s.notes_slide.notes_text_frame.text = "mis notas"
    s.shapes.add_picture(str(tmp_path / "p.png"), 0, 0)
    prs.save(tmp_path / "t.pptx")
    out = tmp_path / "out"
    assert run_main(main, [tmp_path / "t.pptx", out], monkeypatch) == 0
    d = json.loads((out / "extracted.json").read_text(encoding="utf-8"))
    assert d["source"] == "t.pptx" and d["count"] == 1
    sl = d["slides"][0]
    assert sl["title"] == "Hola" and sl["notes"] == "mis notas"
    assert sl["blocks"] == [[{"text": "punto uno", "level": 0}]]
    assert sl["images"][0]["file"] == "assets/slide01-img1.png"
    assert (out / sl["images"][0]["file"]).is_file()
    assert "1 slides" in capsys.readouterr().out
