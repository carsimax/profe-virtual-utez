import json

import pytest
from docx import Document

from cmx_teach.docs.fill import main
from conftest import run_main


@pytest.fixture
def tpl(tmp_path):
    d = Document()
    d.add_paragraph("Hola {{nombre}}, grupo {{grupo}}")
    p = tmp_path / "plantilla.docx"
    d.save(p)
    return p


def data(tmp_path, obj):
    f = tmp_path / "datos.json"
    f.write_text(json.dumps(obj), encoding="utf-8")
    return f


def test_vars_lists_variables(tpl, monkeypatch, capsys):
    assert run_main(main, [tpl, "--vars"], monkeypatch) == 0
    assert capsys.readouterr().out.split() == ["grupo", "nombre"]


def test_fill(tpl, tmp_path, monkeypatch):
    out = tmp_path / "out.docx"
    assert run_main(main, [tpl, data(tmp_path, {"nombre": "Ana", "grupo": "3A"}), out], monkeypatch) == 0
    assert "Hola Ana, grupo 3A" in "\n".join(p.text for p in Document(out).paragraphs)


def test_missing_value_fails(tpl, tmp_path, monkeypatch, capsys):
    out = tmp_path / "out.docx"
    assert run_main(main, [tpl, data(tmp_path, {"nombre": "Ana"}), out], monkeypatch) not in (0, None)
    assert not out.exists()


def test_leftover_tag_fails(tmp_path, monkeypatch):
    d = Document()
    d.add_paragraph("Hola {{nombre}} {{lit}}")
    p = tmp_path / "t.docx"
    d.save(p)
    out = tmp_path / "out.docx"
    assert run_main(main, [p, data(tmp_path, {"nombre": "Ana", "lit": "{{oops}}"}), out], monkeypatch) not in (0, None)


def test_no_args_exit_2(monkeypatch):
    assert run_main(main, [], monkeypatch) == 2
